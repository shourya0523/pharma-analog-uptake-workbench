"""The tracker export folds the verification app's verdicts in correctly.

Builds the real workbook from current gold with a synthetic verdict snapshot
and reads it back: each product-level outcome, the per-row columns, and the
sheet listing every verdict.
"""

from __future__ import annotations

import csv
import importlib.util
import json
import sys
from collections import defaultdict

import pytest
from openpyxl import load_workbook

from tests.test_gold_is_not_an_input import REPO

BUILDER = REPO / "scripts" / "sourcing" / "build_verification_tracker.py"


def _load(path, name):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    tracker = _load(BUILDER, "build_verification_tracker")
    gold = tracker.sourced_rows()
    by_drug = defaultdict(list)
    for gold_id, row in gold.items():
        by_drug[row["drug_name"]].append(gold_id)
    requested = [r["brand_name"] for r in csv.DictReader(tracker.TARGETS.open(newline=""))]
    # Four requested products with quarterly rows, one per outcome.
    drugs = [d for d in requested if len(by_drug.get(d, [])) >= 3][:4]
    done, flagged, needs_fix, partial = drugs
    verdicts, resolutions = [], []

    def verdict(gold_id, reviewer, kind="confirmed", seen=None):
        verdicts.append({"gold_id": gold_id, "reviewer": reviewer, "verdict": kind, "value_seen": None,
                         "gold_value_seen": gold[gold_id]["value_reported"] if seen is None else seen,
                         "note": "", "updated_at": "2026-10-05T09:00:00+00:00"})

    for gold_id in by_drug[done]:
        verdict(gold_id, "a@team.test")
    first_flag = sorted(by_drug[flagged])[0]
    for gold_id in by_drug[flagged]:
        verdict(gold_id, "a@team.test", "wrong_value" if gold_id == first_flag else "confirmed")
    fix = sorted(by_drug[needs_fix])[0]
    for gold_id in by_drug[needs_fix]:
        verdict(gold_id, "b@team.test", "wrong_scope" if gold_id == fix else "confirmed")
    resolutions.append({"gold_id": fix, "outcome": "gold_needs_fix", "note": "line is U.S. only",
                        "resolved_by": "a@team.test", "resolved_at": None})
    stale_row = sorted(by_drug[partial])[0]
    verdict(stale_row, "b@team.test", seen=-1.0)

    snapshot = tmp_path_factory.mktemp("tracker") / "human_verdicts.json"
    snapshot.write_text(json.dumps({"exported_at": "2026-10-05T10:00:00+00:00", "source": "test",
                                    "verdicts": verdicts, "resolutions": resolutions,
                                    "team": {"a@team.test": "Asha", "b@team.test": "Ben"}}))
    out = snapshot.with_name("tracker.xlsx")
    argv = sys.argv
    sys.argv = ["build_verification_tracker.py", "--verdicts", str(snapshot), "--out", str(out)]
    try:
        assert tracker.main() == 0
    finally:
        sys.argv = argv
    book = load_workbook(out)
    return {"book": book, "done": done, "flagged": flagged, "needs_fix": needs_fix, "partial": partial,
            "first_flag": first_flag, "verdicts": verdicts, "by_drug": by_drug, "stale_row": stale_row}


def _sheet_rows(sheet):
    rows = list(sheet.iter_rows(values_only=True))
    return [dict(zip(rows[0], r)) for r in rows[1:]]


def _product(built, drug):
    return next(r for r in _sheet_rows(built["book"]["Drug Checklist"]) if r["Brand Name"] == drug)


def test_a_fully_confirmed_product_is_verified(built):
    row = _product(built, built["done"])
    assert row["Manually Verified"] == "Yes"
    assert row["Verified By"] == "Asha"
    assert row["Date Verified"] == "2026-10-05"
    n = len(built["by_drug"][built["done"]])
    assert row["App Review (rows reviewed / total)"] == f"{n} / {n}"


def test_an_open_flag_or_a_needed_fix_is_not_verified(built):
    flagged = _product(built, built["flagged"])
    assert flagged["Manually Verified"] == "No"
    assert "open flag" in flagged["Reviewer Notes"]
    fix = _product(built, built["needs_fix"])
    assert fix["Manually Verified"] == "No"
    assert "gold needs a fix" in fix["Reviewer Notes"] and "line is U.S. only" in fix["Reviewer Notes"]


def test_partial_review_is_blank_and_a_changed_figure_is_called_out(built):
    row = _product(built, built["partial"])
    assert row["Manually Verified"] in (None, "")
    assert "gold figure changed" in row["Reviewer Notes"]


def test_rows_sheet_carries_each_rows_verdicts(built):
    rows = {r["Gold ID"]: r for r in _sheet_rows(built["book"]["Rows"])}
    assert rows[built["first_flag"]]["Spot-Check (reviewer)"] == "No"
    assert "Asha: wrong value" in rows[built["first_flag"]]["Reviewer Verdicts"]
    assert all(rows[g]["Spot-Check (reviewer)"] == "Yes" for g in built["by_drug"][built["done"]] if g in rows)


def test_every_verdict_is_listed(built):
    listed = _sheet_rows(built["book"]["Human Verdicts"])
    assert len(listed) == len(built["verdicts"])
    stale = [r for r in listed if r["Gold ID"] == built["stale_row"]]
    assert stale and stale[0]["Gold Changed Since"] == "yes"
