"""The file called quarterly_revenue.csv holds published quarters, and says so.

Two defects in one line of code. `export_powerbi_csvs` wrote every datapoint
a run held into a file named for quarterly revenue, with no `period_type`
column - so a cumulative figure filed under a quarter's label was
indistinguishable from the quarter, and anyone charting by `period` built a
curve partly out of year-to-date figures. The columns it wrote were a literal,
as was the workbook's own header list beside it, and between them they dropped
`period_type`, `reported_as` and every other column added to the model after
the lists were typed.

Both now come from `DatapointORM.__table__.columns`, so a column added to the
model reaches both sheets without either list being edited. The filtered file
loses nothing, because all_datapoints.csv beside it is unfiltered.
"""

from __future__ import annotations

import csv
import io

from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from app.db.migrations import upgrade_database
from app.db.models import DatapointORM, DrugJobORM, ExportORM, ExtractionRunORM
from app.domain.models import PeriodType, ValidationStatus
from app.export.builder import DATAPOINT_COLUMNS, QUARTERLY_HEADERS, ExportBuilder
from app.storage.filestore import LocalFileStore

# One of each shape the unfiltered file has to keep and the quarterly one
# has to drop: the published quarter, a cumulative figure wearing a
# quarter's label, and a quarter still held for review.
ROWS = [
    ("dp-q", "2024Q2", PeriodType.QUARTERLY.value, ValidationStatus.AUTO_PASS.value, 12.0),
    ("dp-ytd", "2024Q2", PeriodType.YTD.value, ValidationStatus.AUTO_PASS.value, 21.5),
    ("dp-held", "2024Q3", PeriodType.QUARTERLY.value, ValidationStatus.NEEDS_REVIEW.value, 13.0),
    ("dp-conf", "2024Q4", PeriodType.QUARTERLY.value, ValidationStatus.CONFIRMED.value, 14.0),
]


def _database():
    engine = create_engine("sqlite://")
    upgrade_database(engine)
    db = Session(engine)
    db.add(ExtractionRunORM(id="run", status="completed"))
    db.add(DrugJobORM(id="job", run_id="run", drug_name="Calderon", status="completed"))
    for row_id, period, period_type, status, value in ROWS:
        db.add(
            DatapointORM(
                id=row_id,
                job_id="job",
                period=period,
                period_type=period_type,
                value_normalized_usd_millions=value,
                revenue_scope="Product family",
                reported_as="Calderon and Calderon XR, together",
                source_url="https://sec.gov/a-filing",
                source_quote="Calderon net sales",
                validation_status=status,
                issue_flags=["a_flag", "another"],
            )
        )
    db.commit()
    return db


def test_the_headers_are_the_datapoint_s_own_columns():
    """Every column but the ids and the citation record, in the model's order."""
    described = [
        column.name
        for column in DatapointORM.__table__.columns
        if column.name != "id"
        and not column.name.endswith("_id")
        and column.name != "citation_json"
    ]
    assert DATAPOINT_COLUMNS == described
    assert QUARTERLY_HEADERS == ["drug_name", *described]
    # The three the hand-written list dropped, named because each is the
    # difference between a row that says what it is and one that does not.
    for column in ("period_type", "reported_as", "formulation"):
        assert column in QUARTERLY_HEADERS


async def _csv(store, exports, fmt):
    export = next(item for item in exports if item.format == fmt)
    data = (await store.get(export.storage_key)).decode()
    return list(csv.DictReader(io.StringIO(data)))


async def test_the_quarterly_file_holds_published_quarters_and_the_other_holds_everything(tmp_path):
    db = _database()
    store = LocalFileStore(str(tmp_path))
    exports = await ExportBuilder(db, store).export_powerbi_csvs("run")

    quarterly = await _csv(store, exports, "quarterly_revenue_csv")
    everything = await _csv(store, exports, "all_datapoints_csv")

    assert {row["period_type"] for row in quarterly} == {PeriodType.QUARTERLY.value}
    assert {row["validation_status"] for row in quarterly} == {
        ValidationStatus.AUTO_PASS.value,
        ValidationStatus.CONFIRMED.value,
    }
    assert [row["period"] for row in quarterly] == ["2024Q2", "2024Q4"]
    # Nothing is lost: the rows the filter drops are in the file beside it.
    assert len(everything) == len(ROWS)
    assert {row["period_type"] for row in everything} == {
        PeriodType.QUARTERLY.value,
        PeriodType.YTD.value,
    }


async def test_what_a_figure_was_reported_as_reaches_both_files(tmp_path):
    """A figure a filer prints for two products is not one product's figure."""
    db = _database()
    store = LocalFileStore(str(tmp_path))
    exports = await ExportBuilder(db, store).export_powerbi_csvs("run")

    for fmt in ("quarterly_revenue_csv", "all_datapoints_csv"):
        rows = await _csv(store, exports, fmt)
        assert all(row["reported_as"] == "Calderon and Calderon XR, together" for row in rows)
        assert all(row["issue_flags"] == "a_flag,another" for row in rows)


async def test_the_workbook_sheet_and_the_csv_carry_the_same_columns(tmp_path):
    from openpyxl import load_workbook

    db = _database()
    store = LocalFileStore(str(tmp_path))
    export = await ExportBuilder(db, store).export_product_workbook("job")
    book = load_workbook(io.BytesIO(await store.get(export.storage_key)))
    sheet = book["Quarterly Revenue"]
    header = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
    assert header == QUARTERLY_HEADERS
    # The workbook is the audit copy, so it keeps the held rows too.
    assert sheet.max_row == len(ROWS) + 1


async def test_both_files_are_registered_as_exports(tmp_path):
    db = _database()
    store = LocalFileStore(str(tmp_path))
    exports = await ExportBuilder(db, store).export_powerbi_csvs("run")
    formats = {item.format for item in exports}
    assert {"products_csv", "quarterly_revenue_csv", "all_datapoints_csv"} <= formats
    assert db.query(ExportORM).filter_by(format="all_datapoints_csv").count() == 1
