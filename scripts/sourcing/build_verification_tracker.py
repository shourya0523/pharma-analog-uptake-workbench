"""Write exports/gold_verification_tracker.xlsx: the checklist a reviewer ticks.

One row per requested product (docs/sourcing/target_products.csv), joined to
what gold holds for it and to the verification log, plus a sheet of every
quarterly row with a clickable source so a reviewer can spot-check any of
them. The reviewer columns are left blank and carry a Yes/No dropdown.

The verification log (docs/sourcing/verification_log.json) is keyed by
drug name and records, per product, the automated check of each row and the
verifier's own reading:

    {"Calderon": {"status": "verified", "notes": "...",
                  "automated": {"2024Q1": true, ...},
                  "hand_checked": ["2024Q1", ...]}}

The reviewer columns are filled from the verification app's verdicts when a
snapshot exists (docs/sourcing/human_verdicts.json, written by
pull_verdicts.py); without one they are left blank for hand entry.

    python scripts/sourcing/build_verification_tracker.py [--verdicts PATH] [--out PATH]
"""

from __future__ import annotations

import argparse
import csv
import json
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

REPO_ROOT = Path(__file__).resolve().parents[2]
GOLD = REPO_ROOT / "seed" / "gold"
TARGETS = REPO_ROOT / "docs" / "sourcing" / "target_products.csv"
LOG = REPO_ROOT / "docs" / "sourcing" / "verification_log.json"
OUT = REPO_ROOT / "exports" / "gold_verification_tracker.xlsx"
VERDICTS = REPO_ROOT / "docs" / "sourcing" / "human_verdicts.json"

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
REVIEW_FILL = PatternFill("solid", fgColor="FFF2CC")
STATUS_FILLS = {
    "verified": PatternFill("solid", fgColor="E2EFDA"),
    "verified_with_notes": PatternFill("solid", fgColor="FFF2CC"),
    "excluded_verified": PatternFill("solid", fgColor="DDEBF7"),
    "pending": PatternFill("solid", fgColor="F2F2F2"),
    "failed": PatternFill("solid", fgColor="F8CBAD"),
}


def load_jsonl(name: str) -> list[dict]:
    path = GOLD / name
    return [json.loads(line) for line in path.read_text().splitlines() if line.strip()]


def names_for(brand: str) -> list[str]:
    """A brand cell can name a product family ("Calderon / Calderon XR")."""
    return [brand] + [part.strip() for part in brand.split("/") if part.strip() != brand]


def sourced_rows() -> dict[str, dict]:
    """Every gold row a reviewer can check, by gold_id: any row in any
    seed/gold/*.jsonl with a gold_id and a source_url (the verification app's rule)."""
    out = {}
    for path in sorted(GOLD.glob("*.jsonl")):
        for row in load_jsonl(path.name):
            if row.get("gold_id") and row.get("source_url"):
                out[row["gold_id"]] = row
    return out


class Reviews:
    """The app's verdicts and resolutions, rolled up per row and per product."""

    def __init__(self, snapshot: dict | None, gold: dict[str, dict]):
        self.present = snapshot is not None
        snapshot = snapshot or {"verdicts": [], "resolutions": [], "team": {}}
        self.exported_at = snapshot.get("exported_at", "")
        self.team = {k.lower(): v for k, v in (snapshot.get("team") or {}).items()}
        self.gold = gold
        self.by_row: dict[str, list[dict]] = defaultdict(list)
        for verdict in snapshot["verdicts"]:
            self.by_row[verdict["gold_id"]].append(verdict)
        self.resolution = {r["gold_id"]: r for r in snapshot["resolutions"]}
        self.ids_by_drug: dict[str, list[str]] = defaultdict(list)
        for gold_id, row in gold.items():
            self.ids_by_drug[row["drug_name"]].append(gold_id)

    def name(self, email: str) -> str:
        return self.team.get((email or "").lower()) or (email or "").split("@")[0]

    def status(self, gold_id: str) -> str:
        """unverified / verified / flagged, as the app's row_status view decides it."""
        verdicts = self.by_row.get(gold_id, [])
        if not verdicts:
            return "unverified"
        if any(v["verdict"] != "confirmed" for v in verdicts) and gold_id not in self.resolution:
            return "flagged"
        return "verified"

    def stale(self, verdict: dict) -> bool:
        """The gold figure changed after the reviewer gave this verdict."""
        now = self.gold.get(verdict["gold_id"], {}).get("value_reported")
        seen = verdict.get("gold_value_seen")
        if now is None or seen is None:
            return now is not seen
        return abs(float(now) - float(seen)) > 1e-9

    def row_summary(self, gold_id: str) -> str:
        return "; ".join(
            f"{self.name(v['reviewer'])}: {v['verdict'].replace('_', ' ')}"
            + (f" (read {v['value_seen']:g})" if v.get("value_seen") is not None else "")
            + (" [gold since changed]" if self.stale(v) else "")
            + (f" - {v['note']}" if v.get("note") else "")
            for v in self.by_row.get(gold_id, []))

    def resolution_text(self, gold_id: str) -> str:
        r = self.resolution.get(gold_id)
        if not r:
            return ""
        return f"{r['outcome'].replace('_', ' ')} ({self.name(r.get('resolved_by'))})" + (f": {r['note']}" if r.get("note") else "")

    def product(self, drug: str) -> dict:
        ids = self.ids_by_drug.get(drug, [])
        statuses = {gold_id: self.status(gold_id) for gold_id in ids}
        reviewed = [g for g, s in statuses.items() if s != "unverified"]
        flagged = [g for g, s in statuses.items() if s == "flagged"]
        needs_fix = [g for g in ids if self.resolution.get(g, {}).get("outcome") == "gold_needs_fix"]
        verdicts = [v for g in ids for v in self.by_row.get(g, [])]
        if not verdicts:
            verified = ""
        elif flagged or needs_fix:
            verified = "No"
        elif len(reviewed) == len(ids):
            verified = "Yes"
        else:
            verified = ""
        notes = [f"{len(reviewed)} of {len(ids)} rows reviewed in the app"] if verdicts else []
        notes += [f"open flag {self.gold[g]['period'] or g}: {self.row_summary(g)}" for g in sorted(flagged)]
        notes += [f"gold needs a fix {self.gold[g]['period'] or g}: {self.resolution_text(g)}" for g in sorted(needs_fix)]
        stale = sum(self.stale(v) for v in verdicts)
        if stale:
            notes.append(f"{stale} verdicts given before the gold figure changed")
        return {
            "verified": verified,
            "by": ", ".join(sorted({self.name(v["reviewer"]) for v in verdicts})),
            "date": max((v.get("updated_at") or "" for v in verdicts), default="")[:10],
            "notes": "; ".join(notes),
            "progress": f"{len(reviewed)} / {len(ids)}" if ids else "",
        }


def header(sheet, columns: list[str], review: set[str]) -> None:
    sheet.append(columns)
    for index, name in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index)
        cell.font = Font(bold=True, color="FFFFFF" if name not in review else "000000")
        cell.fill = REVIEW_FILL if name in review else HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = "E2"


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--verdicts", type=Path, default=VERDICTS,
                        help="snapshot from pull_verdicts.py (default: %(default)s, if present)")
    parser.add_argument("--out", type=Path, default=OUT)
    args = parser.parse_args()
    reviews = Reviews(json.loads(args.verdicts.read_text()) if args.verdicts.is_file() else None, sourced_rows())

    targets = list(csv.DictReader(TARGETS.open(newline="")))
    log = json.loads(LOG.read_text()) if LOG.is_file() else {}
    quarterly = load_jsonl("quarterly_revenue.jsonl")
    coverage = {row["drug_name"]: row for row in load_jsonl("series_coverage.jsonl")}
    excluded = {row["drug_name"]: row for row in load_jsonl("excluded_products.jsonl")}
    annual = defaultdict(list)
    for row in load_jsonl("annual_revenue.jsonl"):
        annual[row["drug_name"]].append(row)
    by_drug = defaultdict(list)
    for row in quarterly:
        by_drug[row["drug_name"]].append(row)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Drug Checklist"
    review_columns = ["Manually Verified", "Verified By", "Date Verified", "Reviewer Notes"]
    columns = [
        "ID", "Therapeutic Area", "Brand Name", "Generic / INN", "Gold Status",
        "Issuer", "Series (benchmark identity)", "Revenue Scope", "Currency",
        "First Quarter", "Last Quarter", "Quarters", "Distinct Sources",
        "Derived Rows", "Automated Check (rows passing / total)",
        "Claude Verification", "Claude Notes", "App Review (rows reviewed / total)",
    ] + review_columns
    header(sheet, columns, set(review_columns))

    detail_rows = []
    for target in targets:
        drug = next((n for n in names_for(target["brand_name"]) if n in by_drug or n in excluded or n in annual), target["brand_name"])
        rows = sorted(by_drug.get(drug, []), key=lambda r: r["period"])
        entry = log.get(drug, {})
        automated = entry.get("automated", {})
        issuers = list(dict.fromkeys(r["manufacturer"] for r in rows))
        roles = {r["series_role"] for r in annual.get(drug, [])}
        if rows:
            status = "Quarterly series" + (f" (joined across {len(issuers)} owners)" if len(issuers) > 1 else "")
        elif "peak_benchmark" in roles:
            status = "Annual benchmark (no quarterly figures)"
        elif roles and drug in excluded:
            status = f"Annual context only; excluded ({excluded[drug]['reason_code']})"
        elif drug in excluded:
            status = f"Excluded ({excluded[drug]['reason_code']})"
        else:
            status = "Not yet in gold"
        cov = coverage.get(drug, {})
        review = reviews.product(drug)
        sheet.append([
            target["id"], target["therapeutic_area"], target["brand_name"], target["generic_name"],
            status,
            " -> ".join(issuers),
            rows[0]["benchmark_identity"] if rows else "",
            rows[0]["revenue_scope"] if rows else "",
            rows[0]["currency"] if rows else "",
            rows[0]["period"] if rows else "",
            rows[-1]["period"] if rows else "",
            len(rows) or "",
            len({r["source_url"] for r in rows}) or "",
            sum(not r["derivation"].startswith("direct") for r in rows) if rows else "",
            f"{sum(1 for v in automated.values() if v)} / {len(automated)}" if automated else "",
            entry.get("status", "pending"),
            "; ".join(filter(None, [entry.get("notes", ""), cov.get("series_start_reason", ""), cov.get("series_end_reason", "")])),
            review["progress"],
            review["verified"], review["by"], review["date"], review["notes"],
        ])
        fill = STATUS_FILLS.get(entry.get("status", "pending"))
        if fill:
            sheet.cell(row=sheet.max_row, column=columns.index("Claude Verification") + 1).fill = fill
        for row in rows:
            detail_rows.append((target["id"], row, automated.get(row["period"]), row["period"] in set(entry.get("hand_checked", []))))

    validation = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    sheet.add_data_validation(validation)
    col = get_column_letter(columns.index("Manually Verified") + 1)
    validation.add(f"{col}2:{col}{sheet.max_row}")
    widths = [6, 22, 18, 26, 26, 18, 34, 16, 9, 11, 11, 9, 9, 9, 14, 18, 60, 14, 12, 18, 13, 60]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.auto_filter.ref = sheet.dimensions

    rows_sheet = workbook.create_sheet("Rows")
    row_columns = [
        "ID", "Drug", "Period", "Value Reported (millions)", "Currency",
        "USD Millions", "Derivation", "Source", "Source Quote",
        "Automated Check", "Hand-Checked by Claude", "Spot-Check (reviewer)",
        "Reviewer Verdicts", "Resolution", "Gold ID",
    ]
    review_cols = {"Spot-Check (reviewer)", "Reviewer Verdicts", "Resolution"}
    header(rows_sheet, row_columns, review_cols)
    spot_value = {"verified": "Yes", "flagged": "No", "unverified": ""}
    for target_id, row, auto, hand in detail_rows:
        rows_sheet.append([
            target_id, row["drug_name"], row["period"], row["value_reported"], row["currency"],
            row["value_normalized_usd_millions"], row["derivation"], row["source_url"],
            row["source_quote"],
            "" if auto is None else ("pass" if auto else "FAIL"),
            "yes" if hand else "",
            spot_value[reviews.status(row["gold_id"])],
            reviews.row_summary(row["gold_id"]),
            reviews.resolution_text(row["gold_id"]),
            row["gold_id"],
        ])
        link = rows_sheet.cell(row=rows_sheet.max_row, column=8)
        link.hyperlink = row["source_url"]
        link.font = Font(color="0563C1", underline="single")
    spot = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    rows_sheet.add_data_validation(spot)
    spot.add(f"L2:L{max(rows_sheet.max_row, 2)}")
    for index, width in enumerate([6, 18, 9, 12, 8, 12, 24, 50, 80, 10, 10, 12, 50, 40, 40], start=1):
        rows_sheet.column_dimensions[get_column_letter(index)].width = width
    rows_sheet.auto_filter.ref = rows_sheet.dimensions

    # Every verdict the app holds, on any kind of row (quarterly, annual,
    # companion, exclusion), so nothing reviewed is lost from the export.
    verdict_sheet = workbook.create_sheet("Human Verdicts")
    verdict_columns = ["Gold ID", "Drug", "Period", "Gold Value Now", "Reviewer", "Verdict",
                       "Value Read", "Gold Value When Reviewed", "Gold Changed Since", "Note",
                       "Reviewed At", "Row Status", "Resolution"]
    header(verdict_sheet, verdict_columns, set())
    verdict_sheet.freeze_panes = "B2"
    for gold_id in sorted(reviews.by_row):
        gold_row = reviews.gold.get(gold_id, {})
        for v in reviews.by_row[gold_id]:
            verdict_sheet.append([
                gold_id, gold_row.get("drug_name", "(no longer in gold)"), gold_row.get("period", ""),
                gold_row.get("value_reported"), reviews.name(v["reviewer"]), v["verdict"].replace("_", " "),
                v.get("value_seen"), v.get("gold_value_seen"), "yes" if reviews.stale(v) else "",
                v.get("note", ""), (v.get("updated_at") or "")[:19].replace("T", " "),
                reviews.status(gold_id), reviews.resolution_text(gold_id),
            ])
    for index, width in enumerate([44, 18, 9, 12, 14, 14, 11, 12, 10, 50, 18, 11, 40], start=1):
        verdict_sheet.column_dimensions[get_column_letter(index)].width = width
    verdict_sheet.auto_filter.ref = verdict_sheet.dimensions

    readme = workbook.create_sheet("Read Me", 0)
    for line in [
        ["Gold Verification Tracker"],
        [],
        ["Drug Checklist", "One row per requested product. Yellow columns are yours: set Manually Verified to Yes once you have checked the product."],
        ["Rows", "Every quarterly row behind those products, with a clickable source. Spot-Check is optional per row."],
        [],
        ["Automated Check", "Every row's source was fetched and the row label plus every quoted figure was found on one line of that document, in order."],
        ["Claude Verification", "verified = automated check passed and Claude read the hand-checked rows at source; verified_with_notes = passed with a caveat in Claude Notes; excluded_verified = the exclusion evidence was read at source; pending = not yet reviewed."],
        ["Hand-Checked by Claude", "Rows Claude opened and read in the source document: every derived row, every failing row, each series' first, last and peak quarter, and a sample of the rest."],
        [],
        ["Reviewer columns", (f"Filled from the verification app's verdicts as of {reviews.exported_at}. " if reviews.present else "No verdict snapshot was found; fill these by hand. ")
         + "Manually Verified is Yes when every gold row of the product has a verdict and none is an open flag or a resolved 'gold needs a fix'; No when one is; blank while review is in progress."],
        ["Human Verdicts", "Every verdict from the app, on any kind of gold row, with whether the gold figure has changed since it was given."],
    ]:
        readme.append(line)
    readme["A1"].font = Font(bold=True, size=14)
    readme.column_dimensions["A"].width = 24
    readme.column_dimensions["B"].width = 120

    args.out.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.out)
    print(f"wrote {args.out} ({len(targets)} products, {len(detail_rows)} rows, "
          f"{sum(map(len, reviews.by_row.values()))} verdicts"
          + (f" as of {reviews.exported_at})" if reviews.present else ", no verdict snapshot)"))
    return 0


if __name__ == "__main__":
    sys.exit(main())
