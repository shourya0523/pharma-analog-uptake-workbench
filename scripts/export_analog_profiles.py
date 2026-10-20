"""Write the analog profiles workbook in the four-section layout of the rubric.

Sections, their descriptions, criteria and options are read from
docs/analog_profiles.md, so the workbook follows the rubric without a second
list to keep in step. Profile values and their evidence are read from
seed/analog_profiles.csv and seed/analog_profile_evidence.csv.

Sheets:
  Criteria          the four sections, each with its description and its
                    criteria and options, as the rubric states them
  Product Profiles  one row per product; a header row names the section each
                    column belongs to. The Revenue section after the product
                    columns is the product's record from
                    export_peak_candidates.peak_records()
  Profile Evidence  one row per value: the document it was read from and the
                    sentence that decides it
  Quarterly Revenue one row per product-quarter in gold, for this workbook's
                    products only, with the rolling four-quarter sum ending there
  Annual Revenue    the peak-benchmark years of products gold holds only annually
  Revenue Matrix    products by calendar quarter: USD millions, then a second
                    block of rolling four-quarter sums

Profiled products gold holds no revenue for keep their row, with the reason in
peak_status, and are printed when the workbook is written.

    python scripts/export_analog_profiles.py [OUT.xlsx] [--profiles P.csv --evidence E.csv]
"""

from __future__ import annotations

import argparse
import csv
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from export_peak_candidates import label, load, peak_records, quarter_index, rolling_windows

REPO = Path(__file__).resolve().parents[1]
RUBRIC = REPO / "docs" / "analog_profiles.md"
PROFILES = REPO / "seed" / "analog_profiles.csv"
EVIDENCE = REPO / "seed" / "analog_profile_evidence.csv"
OUT = REPO / "exports" / "analog_profiles.xlsx"

INK = "1F3864"
SECTION_FILLS = ["2E75B6", "548235", "BF8F00", "7030A0"]
FONT = "Arial"
THIN = Side(style="thin", color="BFBFBF")
GRID = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
REVENUE_FILL = "C55A11"
ID_COLUMNS = ["drug_name", "generic_name", "launch_indication"]
# Fields of a peak_records() row, in the order the Revenue section shows them;
# currencies_reported is added here from the product's gold rows.
REVENUE_COLUMNS = [
    "peak_status", "first_quarter", "last_quarter", "quarters",
    "best_rolling_4q_usd_m", "best_window", "best_window_is_latest",
    "latest_rolling_4q_usd_m", "latest_window", "latest_vs_best",
    "best_annual_usd_m", "peak_year", "currencies_reported",
]
QUARTER_COLUMNS = [
    "drug_name", "period", "value_reported", "currency", "source_value_reported", "source_unit",
    "value_normalized_usd_millions", "rolling_4q_usd_m", "pct_of_best_rolling_4q",
    "issuer", "manufacturer", "revenue_scope", "geography", "source_url", "source_quote",
]
ANNUAL_COLUMNS = [
    "drug_name", "period", "value_reported", "currency", "source_value_reported", "source_unit",
    "value_normalized_usd_millions", "fx_rate_to_usd", "manufacturer", "revenue_scope",
    "geography", "series_role", "source_url", "source_quote",
]
WIDE = {"drug_name": 18, "source_url": 50, "source_quote": 80, "peak_status": 26,
        "revenue_scope": 22, "geography": 18, "manufacturer": 20, "issuer": 18}


def sections() -> list[dict]:
    """[{title, description, criteria: [(name, options)]}] in rubric order."""
    text = RUBRIC.read_text()
    out = []
    for block in re.split(r"^## ", text, flags=re.M)[1:]:
        title, _, body = block.partition("\n")
        rows = re.findall(r"^\| ([^|]+?) \| ([^|]+?) \|", body, re.M)
        criteria = [(a.strip(), b.strip()) for a, b in rows
                    if a.strip() != "Criterion" and not set(a.strip()) <= {"-"}]
        if not criteria:
            continue
        description = next((p.strip() for p in body.split("\n\n") if p.strip() and not p.strip().startswith("|")), "")
        out.append({"title": title.strip(), "description": description, "criteria": criteria})
    return out


def read(path: Path) -> list[dict]:
    if not path.exists():
        return []
    with path.open(newline="") as handle:
        return list(csv.DictReader(handle))


def cell(ws, row, col, value, *, bold=False, color=None, fill=None, wrap=False, size=10, italic=False):
    c = ws.cell(row, col, value)
    c.font = Font(name=FONT, bold=bold, color=color, size=size, italic=italic)
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    c.alignment = Alignment(wrap_text=wrap, vertical="top")
    return c


def number_format(name: str) -> str | None:
    if name.endswith("_usd_m") or name.endswith("_usd_millions"):
        return "#,##0.0"
    if name in ("latest_vs_best", "pct_of_best_rolling_4q"):
        return "0.0%"
    return None


def revenue(names: list[str]) -> dict:
    """Gold revenue for the profiled products, keyed by profile drug_name.

    A profile name is matched to gold's drug_name exactly, then ignoring case.
    Returns {"record": {name: peak record}, "quarters": {name: [rows with
    rolling_4q_usd_m and pct_of_best_rolling_4q]}, "annual": {name: rows},
    "missing": {name: reason}}.
    """
    quarterly = defaultdict(list)
    for row in load("quarterly_revenue.jsonl"):
        quarterly[row["drug_name"]].append(row)
    annual = defaultdict(list)
    for row in load("annual_revenue.jsonl"):
        if row.get("series_role") == "peak_benchmark":
            annual[row["drug_name"]].append(row)
    excluded = {row["drug_name"]: row for row in load("excluded_products.jsonl")}
    records = {row["drug_name"]: row for row in peak_records()}
    known = set(quarterly) | set(annual) | set(records) | set(excluded)
    folded = {name.casefold(): name for name in known}

    out = {"record": {}, "quarters": {}, "annual": {}, "missing": {}}
    for name in names:
        gold = name if name in known else folded.get(name.casefold())
        rows = sorted(quarterly.get(gold, []), key=lambda r: r["period"])
        years = sorted(annual.get(gold, []), key=lambda r: r["period"])
        if gold in records:
            record = dict(records[gold])
            record["currencies_reported"] = ", ".join(sorted({r["currency"] for r in rows or years}))
            out["record"][name] = record
        if rows:
            windows = rolling_windows(rows)
            ending = {last: total for _, last, total in windows}
            best = max((w[2] for w in windows), default=None)
            for row in rows:
                rolled = ending.get(label(quarter_index(row["period"])))
                row = dict(row, rolling_4q_usd_m=rolled,
                           pct_of_best_rolling_4q=rolled / best if rolled is not None and best else None)
                out["quarters"].setdefault(name, []).append(row)
        elif years:
            out["annual"][name] = years
        if name not in out["record"]:
            if gold in excluded:
                reason = f"excluded from gold: {excluded[gold].get('reason_code')}"
            elif years:
                reason = "annual-only, not a peak record"
            else:
                reason = "not in gold"
            out["missing"][name] = reason
    return out


def table(ws, columns, rows, *, top=1):
    """Write a header row at ``top`` and ``rows`` under it; freeze and filter."""
    for col, name in enumerate(columns, start=1):
        cell(ws, top, col, name, bold=True, color="FFFFFF", fill=INK, wrap=True).border = GRID
        ws.column_dimensions[get_column_letter(col)].width = WIDE.get(name, 15)
    for r, record in enumerate(rows, start=top + 1):
        for col, name in enumerate(columns, start=1):
            c = cell(ws, r, col, record.get(name), wrap=name == "source_quote")
            c.border = GRID
            if number_format(name):
                c.number_format = number_format(name)
    ws.freeze_panes = f"C{top + 1}"
    ws.auto_filter.ref = f"A{top}:{get_column_letter(len(columns))}{len(rows) + top}"


def matrix(ws, series: dict[str, list[dict]]) -> None:
    """Products by calendar quarter: USD millions, then rolling 4Q sums below."""
    if not series:
        cell(ws, 1, 1, "No product in this workbook has quarterly revenue in gold.", italic=True)
        return
    indices = [quarter_index(r["period"]) for rows in series.values() for r in rows]
    quarters = [label(i) for i in range(min(indices), max(indices) + 1)]
    names = sorted(series, key=str.lower)
    top = 1
    for title, field in (("Quarterly revenue, USD millions", "value_normalized_usd_millions"),
                         ("Rolling four-quarter revenue, USD millions (blank unless all four quarters are present)",
                          "rolling_4q_usd_m")):
        cell(ws, top, 1, title, bold=True, size=12, color=INK)
        header = top + 1
        cell(ws, header, 1, "drug_name", bold=True, color="FFFFFF", fill=INK).border = GRID
        for col, quarter in enumerate(quarters, start=2):
            cell(ws, header, col, quarter, bold=True, color="FFFFFF", fill=INK).border = GRID
        for r, name in enumerate(names, start=header + 1):
            cell(ws, r, 1, name).border = GRID
            values = {row["period"]: row.get(field) for row in series[name]}
            for col, quarter in enumerate(quarters, start=2):
                c = cell(ws, r, col, values.get(quarter))
                c.border = GRID
                c.number_format = "#,##0.0"
        if top == 1:
            ws.auto_filter.ref = f"A{header}:{get_column_letter(len(quarters) + 1)}{header + len(names)}"
        top = header + len(names) + 2
    ws.column_dimensions["A"].width = 18
    for col in range(2, len(quarters) + 2):
        ws.column_dimensions[get_column_letter(col)].width = 9
    ws.freeze_panes = "B3"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("out", nargs="?", type=Path, default=OUT)
    parser.add_argument("--profiles", type=Path, default=PROFILES)
    parser.add_argument("--evidence", type=Path, default=EVIDENCE)
    args = parser.parse_args()
    parts = sections()
    profiles = sorted(read(args.profiles), key=lambda r: r["drug_name"].lower())
    evidence = read(args.evidence)
    gold = revenue([p["drug_name"] for p in profiles])
    for profile in profiles:
        name = profile["drug_name"]
        profile.update(gold["record"].get(name)
                       or {"peak_status": f"no gold revenue - {gold['missing'][name]}"})

    wb = Workbook()
    ws = wb.active
    ws.title = "Criteria"
    row = 1
    cell(ws, row, 1, "Analog selection criteria", bold=True, size=14, color=INK)
    row += 2
    for index, part in enumerate(parts):
        fill = SECTION_FILLS[index % len(SECTION_FILLS)]
        cell(ws, row, 1, part["title"].title(), bold=True, color="FFFFFF", fill=fill, size=12)
        cell(ws, row, 2, None, fill=fill)
        row += 1
        cell(ws, row, 1, part["description"], wrap=True, italic=True)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        ws.row_dimensions[row].height = 48
        row += 1
        cell(ws, row, 1, "Criterion", bold=True, color="FFFFFF", fill=INK).border = GRID
        cell(ws, row, 2, "Variable options", bold=True, color="FFFFFF", fill=INK).border = GRID
        row += 1
        for name, options in part["criteria"]:
            cell(ws, row, 1, name).border = GRID
            cell(ws, row, 2, options, wrap=True).border = GRID
            row += 1
        row += 1
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 110

    ws = wb.create_sheet("Product Profiles")
    columns = list(ID_COLUMNS) + list(REVENUE_COLUMNS)
    owners = [None] * len(ID_COLUMNS) + ["revenue"] * len(REVENUE_COLUMNS)
    for index, part in enumerate(parts):
        for name, _ in part["criteria"]:
            columns.append(name)
            owners.append(index)
    for col, name in enumerate(columns, start=1):
        owner = owners[col - 1]
        if owner is None:
            fill, title, heading = INK, "Product", name.replace("_", " ").capitalize()
        elif owner == "revenue":
            fill, title, heading = REVENUE_FILL, "Revenue", name
        else:
            fill, title, heading = SECTION_FILLS[owner % len(SECTION_FILLS)], parts[owner]["title"].title(), name
        cell(ws, 1, col, title, bold=True, color="FFFFFF", fill=fill)
        cell(ws, 2, col, heading, bold=True, color="FFFFFF", fill=fill, wrap=True).border = GRID
        ws.column_dimensions[get_column_letter(col)].width = (
            34 if name == "launch_indication" else 26 if name == "peak_status" else 18)
    start = 1
    for col in range(2, len(columns) + 2):
        if col > len(columns) or owners[col - 1] != owners[start - 1]:
            if col - 1 > start:
                ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=col - 1)
            start = col
    ws.row_dimensions[2].height = 42
    for r, profile in enumerate(profiles, start=3):
        for col, name in enumerate(columns, start=1):
            value = profile.get(name)
            if name == "Launch Year" and value and value.isdigit():
                value = int(value)
            c = cell(ws, r, col, value, wrap=(name == "launch_indication"))
            c.border = GRID
            if number_format(name):
                c.number_format = number_format(name)
            if value == "Unknown":
                c.font = Font(name=FONT, size=10, color="808080", italic=True)
    ws.freeze_panes = "B3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(columns))}{len(profiles) + 2}"

    ws = wb.create_sheet("Profile Evidence")
    ev_cols = ["drug_name", "criterion", "value", "basis", "source_url", "source_quote"]
    widths = {"drug_name": 18, "criterion": 28, "value": 24, "basis": 60, "source_url": 50, "source_quote": 80}
    for col, name in enumerate(ev_cols, start=1):
        cell(ws, 1, col, name, bold=True, color="FFFFFF", fill=INK).border = GRID
        ws.column_dimensions[get_column_letter(col)].width = widths[name]
    order = {name: i for i, name in enumerate(columns)}
    evidence.sort(key=lambda r: (r["drug_name"].lower(), order.get(r["criterion"], 99)))
    for r, record in enumerate(evidence, start=2):
        for col, name in enumerate(ev_cols, start=1):
            cell(ws, r, col, record.get(name), wrap=name in ("basis", "source_quote")).border = GRID
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:F{len(evidence) + 1}"

    quarter_rows = [row for name in sorted(gold["quarters"], key=str.lower) for row in gold["quarters"][name]]
    table(wb.create_sheet("Quarterly Revenue"), QUARTER_COLUMNS, quarter_rows)
    annual_rows = [row for name in sorted(gold["annual"], key=str.lower) for row in gold["annual"][name]]
    table(wb.create_sheet("Annual Revenue"), ANNUAL_COLUMNS, annual_rows)
    matrix(wb.create_sheet("Revenue Matrix"), gold["quarters"])

    args.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out)
    print(f"wrote {args.out}: {len(profiles)} products, {len(evidence)} evidence rows, {len(parts)} sections, "
          f"{len(quarter_rows)} quarterly and {len(annual_rows)} annual revenue rows")
    for name, reason in sorted(gold["missing"].items(), key=lambda kv: kv[0].lower()):
        print(f"  no gold revenue: {name} ({reason})")


if __name__ == "__main__":
    main()
